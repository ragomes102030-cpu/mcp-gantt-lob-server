"""Testes de models/gantt.py — lógica de cores por status, marco e
destaque de caminho crítico. Não testa a lib xlsx-gantt em si (é
dependência externa já validada manualmente), só a lógica que decide
COMO usá-la a partir de uma atividade do domínio (planejamento de obras).
"""
from __future__ import annotations

import io
from datetime import date

import openpyxl
import pytest

from mcp_gantt_lob.models.gantt import (
    COR_ATRASADA,
    COR_CONCLUIDA,
    AtividadeInput,
    atividade_from_dict,
    gerar_gantt_xlsx,
    _dt,
    _status_ranges,
)


def _wb(xlsx_bytes: bytes):
    return openpyxl.load_workbook(io.BytesIO(xlsx_bytes))


def test_gerar_gantt_lista_vazia_falha():
    with pytest.raises(ValueError, match="vazia"):
        gerar_gantt_xlsx("Projeto", [])


def test_marco_inicio_igual_fim_concluido_fica_verde():
    a = AtividadeInput(
        nome="Entrega das chaves", secao="Marcos",
        data_inicio=date(2026, 3, 1), data_fim_planejada=date(2026, 3, 1),
        progresso=100,
    )
    ranges = _status_ranges(a, cor_planejado="4472C4")
    assert len(ranges) == 1
    assert ranges[0].color == COR_CONCLUIDA


def test_marco_nao_concluido_usa_cor_planejada():
    a = AtividadeInput(
        nome="Entrega das chaves", secao="Marcos",
        data_inicio=date(2026, 3, 1), data_fim_planejada=date(2026, 3, 1),
        progresso=0,
    )
    ranges = _status_ranges(a, cor_planejado="4472C4")
    assert ranges[0].color == "4472C4"


def test_concluida_no_prazo_fica_verde():
    a = AtividadeInput(
        nome="Escavação", secao="Fundação",
        data_inicio=date(2026, 1, 5), data_fim_planejada=date(2026, 1, 10),
        progresso=100,
    )
    ranges = _status_ranges(a, cor_planejado="4472C4")
    assert len(ranges) == 1
    assert ranges[0].color == COR_CONCLUIDA


def test_concluida_com_atraso_fica_vermelha():
    a = AtividadeInput(
        nome="Escavação", secao="Fundação",
        data_inicio=date(2026, 1, 5), data_fim_planejada=date(2026, 1, 10),
        progresso=100, data_fim_prevista_real=date(2026, 1, 15),
    )
    ranges = _status_ranges(a, cor_planejado="4472C4")
    assert ranges[0].color == COR_ATRASADA
    assert ranges[0].end == _dt(a.data_fim_prevista_real)


def test_em_andamento_com_nova_previsao_de_atraso_gera_dois_trechos():
    a = AtividadeInput(
        nome="Concretagem", secao="Fundação",
        data_inicio=date(2026, 1, 8), data_fim_planejada=date(2026, 1, 12),
        progresso=40, data_fim_prevista_real=date(2026, 1, 18),
    )
    ranges = _status_ranges(a, cor_planejado="4472C4")
    assert len(ranges) == 2
    assert ranges[0].color == "4472C4"  # trecho planejado original
    assert ranges[1].color == COR_ATRASADA  # trecho de atraso
    assert ranges[1].start == _dt(a.data_fim_planejada)
    assert ranges[1].end == _dt(a.data_fim_prevista_real)


def test_em_andamento_passou_do_prazo_sem_nova_previsao_sinaliza_atraso():
    a = AtividadeInput(
        nome="Concretagem", secao="Fundação",
        data_inicio=date(2026, 1, 8), data_fim_planejada=date(2026, 1, 12),
        progresso=40, hoje=date(2026, 1, 20),
    )
    ranges = _status_ranges(a, cor_planejado="4472C4")
    assert len(ranges) == 2
    assert ranges[1].color == COR_ATRASADA
    assert ranges[1].end == _dt(date(2026, 1, 20))


def test_em_andamento_dentro_do_prazo_um_trecho_so():
    a = AtividadeInput(
        nome="Concretagem", secao="Fundação",
        data_inicio=date(2026, 1, 8), data_fim_planejada=date(2026, 1, 12),
        progresso=40, hoje=date(2026, 1, 10),
    )
    ranges = _status_ranges(a, cor_planejado="4472C4")
    assert len(ranges) == 1


def test_atividade_from_dict_campos_obrigatorios_e_opcionais():
    a = atividade_from_dict({
        "nome": "Alvenaria", "secao": "Estrutura",
        "data_inicio": "2026-01-13", "data_fim_planejada": "2026-01-22",
        "progresso": 50, "critico": True,
    })
    assert a.nome == "Alvenaria"
    assert a.data_inicio == date(2026, 1, 13)
    assert a.critico is True
    assert a.data_fim_prevista_real is None


def test_atividade_from_dict_critico_default_falso():
    a = atividade_from_dict({
        "nome": "X", "secao": "Y",
        "data_inicio": "2026-01-01", "data_fim_planejada": "2026-01-02",
    })
    assert a.critico is False


def test_gerar_gantt_marca_atividade_critica_em_negrito():
    atividades = [
        AtividadeInput(nome="Escavação", secao="Fundação",
                       data_inicio=date(2026, 1, 5), data_fim_planejada=date(2026, 1, 7),
                       progresso=100, critico=True),
        AtividadeInput(nome="Concretagem", secao="Fundação",
                       data_inicio=date(2026, 1, 8), data_fim_planejada=date(2026, 1, 12),
                       progresso=40, critico=False),
    ]
    wb = _wb(gerar_gantt_xlsx("Teste", atividades, tema="ocean"))
    ws = wb.active
    # linha 4 = primeira atividade (Escavação, crítica); linha 5 = Concretagem
    assert ws.cell(row=4, column=2).font.bold is True
    assert ws.cell(row=5, column=2).font.bold is not True


def test_gerar_gantt_agrupa_por_secao_preservando_ordem():
    atividades = [
        AtividadeInput(nome="A", secao="Fundação", data_inicio=date(2026, 1, 1), data_fim_planejada=date(2026, 1, 2), progresso=0),
        AtividadeInput(nome="B", secao="Estrutura", data_inicio=date(2026, 1, 3), data_fim_planejada=date(2026, 1, 4), progresso=0),
        AtividadeInput(nome="C", secao="Fundação", data_inicio=date(2026, 1, 5), data_fim_planejada=date(2026, 1, 6), progresso=0),
    ]
    wb = _wb(gerar_gantt_xlsx("Teste", atividades, tema="ocean"))
    ws = wb.active
    # Ordem de ENTRADA das seções é preservada (Fundação antes de Estrutura),
    # mesmo com atividades intercaladas — "C" (Fundação) fica junto de "A".
    nomes_col2 = [ws.cell(row=r, column=2).value for r in (4, 5, 6)]
    assert nomes_col2 == ["A", "C", "B"]
