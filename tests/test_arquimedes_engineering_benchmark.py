"""Arquimedes Engineering Benchmark — independent LOB/Gantt reference fixtures.

Dates are deliberately fixed; no production data, network or persistent DB.
"""
from __future__ import annotations

import io
from datetime import date

import openpyxl
import pytest

from mcp_gantt_lob.models.gantt import AtividadeInput, atividade_from_dict, gerar_gantt_xlsx
from mcp_gantt_lob.models.lob import (
    AtividadeLOB, calcular_linha_balanco, dimensionar_equipes
)


def test_lob_three_floors_independent_hand_reference():
    # Foundation A tr=4: 0-4,4-8,8-12.
    # Masonry B tr=6: 4-10,10-16,16-22.
    out = calcular_linha_balanco(
        [AtividadeLOB("Fundacao", 4), AtividadeLOB("Alvenaria", 6)],
        ["P01", "P02", "P03"], data_inicio=date(2026, 11, 2),
    )
    by_name = {x["nome"]: x for x in out["atividades"]}
    assert [(u["inicio_dia"], u["fim_dia"]) for u in by_name["Fundacao"]["unidades"]] == [
        (0, 4), (4, 8), (8, 12)]
    assert [(u["inicio_dia"], u["fim_dia"]) for u in by_name["Alvenaria"]["unidades"]] == [
        (4, 10), (10, 16), (16, 22)]
    assert out["duracao_total_dias"] == 22
    assert by_name["Alvenaria"]["risco_interferencia"] is False
    assert by_name["Fundacao"]["unidades"][0]["data_inicio"] == "2026-11-02"


def test_lob_unequal_ritms_no_overlapping_crew():
    out = calcular_linha_balanco(
        [AtividadeLOB("Estrutura", 6), AtividadeLOB("Vedacao", 4)],
        ["P01", "P02", "P03", "P04"],
    )
    by = {x["nome"]: x for x in out["atividades"]}
    assert [u["inicio_dia"] for u in by["Vedacao"]["unidades"]] == [6, 12, 18, 24]
    assert [u["espera_dias"] for u in by["Vedacao"]["unidades"]] == [0, 2, 2, 2]
    assert by["Vedacao"]["risco_interferencia"] is True
    for a in by.values():
        for older, newer in zip(a["unidades"], a["unidades"][1:]):
            assert newer["inicio_dia"] >= older["fim_dia"]
    assert out["duracao_total_dias"] == 28


def test_lob_fractional_days_have_precise_datetime():
    out = calcular_linha_balanco(
        [AtividadeLOB("Pintura", 1, numero_equipes=2)],
        ["P01", "P02"], data_inicio=date(2026, 11, 2)
    )
    un = out["atividades"][0]["unidades"]
    assert un[0]["data_hora_inicio"] == "2026-11-02T00:00:00"
    assert un[0]["data_hora_fim"] == "2026-11-02T12:00:00"
    assert un[1]["data_hora_inicio"] == "2026-11-02T12:00:00"
    assert un[1]["data_hora_fim"] == "2026-11-03T00:00:00"
    assert "sem calendario" in out["convencao_temporal"]


@pytest.mark.parametrize("invalid", [float("nan"), float("inf"), float("-inf"), 0, -1])
def test_lob_rejects_bad_times(invalid):
    with pytest.raises(ValueError):
        AtividadeLOB("Revestimento", invalid)
    with pytest.raises(ValueError):
        dimensionar_equipes(invalid, 2)


@pytest.mark.parametrize("progress", [-0.01, 100.01, float("nan"), float("inf")])
def test_gantt_rejects_bad_progress(progress):
    with pytest.raises(ValueError, match="progresso"):
        atividade_from_dict({
            "nome": "Alvenaria", "secao": "Ficticia",
            "data_inicio": "2026-11-02", "data_fim_planejada": "2026-11-03",
            "progresso": progress,
        })


def test_gantt_rejects_reversed_dates_before_export():
    with pytest.raises(ValueError, match="data final"):
        atividade_from_dict({
            "nome": "Alvenaria", "secao": "Ficticia",
            "data_inicio": "2026-11-03", "data_fim_planejada": "2026-11-02",
        })


def test_xlsx_file_integrity_and_critical_marker_from_input():
    items = [
        AtividadeInput("Fundacao", "Estrutura", date(2026,11,2),
                       date(2026,11,6), 50, critico=True),
        AtividadeInput("Alvenaria", "Estrutura", date(2026,11,6),
                       date(2026,11,12), 0, critico=False),
    ]
    payload = gerar_gantt_xlsx("Obra Ficticia", items)
    assert payload[:2] == b"PK"  # Valid XLSX/OOXML ZIP header
    wb = openpyxl.load_workbook(io.BytesIO(payload), read_only=False)
    ws = wb.active
    assert [ws.cell(row=i, column=2).value for i in (4,5)] == ["Fundacao", "Alvenaria"]
    assert ws.cell(row=4, column=2).font.bold is True
    assert ws.cell(row=5, column=2).font.bold is not True
    assert ws.row_dimensions[4].outlineLevel == 1
